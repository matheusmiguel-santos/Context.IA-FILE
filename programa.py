import os
from docx import Document
from fpdf import FPDF

def txt_to_pdf(filename, content):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(0, 10, txt=content, ln=True)
    pdf.output(filename)

def txt_to_docx(filename, content):
    doc = Document()
    doc.add_paragraph(content)
    doc.save(filename)

def txt_to_txt(filename, content):
    with open(filename, 'w') as f:
        f.write(content)

def main():
    directory = input('Por favor, digite o caminho do diretório que contém os arquivos .txt: ')
    output_format = input('Por favor, digite o formato de saída desejado (.pdf, .docx, .txt): ')
    output_filename = 'output' + output_format

    content = ''
    for filename in os.listdir(directory):
        if filename.endswith('.txt'):
            with open(os.path.join(directory, filename), 'r') as f:
                content += f.read() + '\n'

    if output_format == '.pdf':
        txt_to_pdf(output_filename, content)
    elif output_format == '.docx':
        txt_to_docx(output_filename, content)
    elif output_format == '.txt':
        txt_to_txt(output_filename, content)
    else:
        print('Formato de arquivo desconhecido.')

if __name__ == '__main__':
    main()
